from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
import zipfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PYTHON = Path(sys.executable)
sys.path.insert(0, str(ROOT / "src"))

from screen_mouse_recorder.app import ScreenMouseRecorderApp
from screen_mouse_recorder.analysis import generate_behavior_report
from screen_mouse_recorder.calibration import build_calibration_result
from screen_mouse_recorder.config import AppConfig
from screen_mouse_recorder.models import Region, TimingContext
from screen_mouse_recorder.mouse_logger import MouseActivityLogger
from screen_mouse_recorder.postprocess import generate_summary
from screen_mouse_recorder.storage import JsonlWriter, SessionStorage
from screen_mouse_recorder.video_recorder import FFmpegRecorder


class CoreSmokeTests(unittest.TestCase):
    def test_region_mapping(self) -> None:
        region = Region(screen_x=100, screen_y=200, width=400, height=300)
        mapped = region.map_point(300, 350)
        self.assertEqual(mapped["region_x"], 200)
        self.assertEqual(mapped["region_y"], 150)
        self.assertEqual(mapped["region_x_norm"], 0.5)
        self.assertEqual(mapped["region_y_norm"], 0.5)
        self.assertTrue(mapped["inside_region"])

    def test_region_even_sized_for_h264(self) -> None:
        region = Region(screen_x=-1544, screen_y=70, width=951, height=783)
        normalized = region.even_sized()

        self.assertEqual(normalized.width, 950)
        self.assertEqual(normalized.height, 782)
        self.assertEqual(normalized.screen_x, region.screen_x)
        self.assertEqual(normalized.screen_y, region.screen_y)

    def test_paused_duration_is_removed_from_video_time(self) -> None:
        timing = TimingContext(
            session_id="test",
            logger_start_monotonic_ms=1000,
            video_zero_monotonic_ms=1200,
            paused_duration_ms=500,
        )
        self.assertEqual(timing.t_video_ms(2200), 500)
        self.assertEqual(timing.timing_dict()["paused_duration_ms"], 500)

    def test_session_name_sanitizer(self) -> None:
        self.assertEqual(ScreenMouseRecorderApp._sanitize_session_name(" demo / round:1 "), "demo_round_1")
        self.assertEqual(ScreenMouseRecorderApp._sanitize_session_name("中文 任务"), "中文_任务")

    def test_config_includes_recording_status_banner_toggle(self) -> None:
        config = AppConfig(show_recording_status_banner=False)

        self.assertFalse(config.show_recording_status_banner)
        self.assertFalse(config.to_dict()["show_recording_status_banner"])

    def test_video_recorder_draws_mouse_cursor(self) -> None:
        region = Region(screen_x=10, screen_y=20, width=320, height=240)
        command = FFmpegRecorder._build_command("ffmpeg", region, Path("recording.mp4"), 30)
        draw_mouse_index = command.index("-draw_mouse")

        self.assertEqual(command[draw_mouse_index + 1], "1")

    def test_calibration_rejects_obvious_misclick(self) -> None:
        region = Region(screen_x=100, screen_y=200, width=400, height=300)
        clicks = self._calibration_clicks(region, offset_x=120, offset_y=0, event_offset_x=120, event_offset_y=0)

        result = build_calibration_result(region, clicks, click_tolerance_px=80)

        self.assertFalse(result["completed"])
        self.assertIn("可能点错检查区域", "；".join(result["failure_reasons"]))

    def test_video_coordinates_stay_raw_with_coordinate_check_data(self) -> None:
        with TemporaryDirectory() as directory:
            region = Region(screen_x=100, screen_y=200, width=400, height=300)
            calibration = build_calibration_result(region, self._calibration_clicks(region, offset_x=4, offset_y=-2), 80)
            timing = TimingContext(session_id="test", logger_start_monotonic_ms=1000, video_zero_monotonic_ms=1000)
            logger = MouseActivityLogger(
                region,
                timing,
                AppConfig(),
                JsonlWriter(Path(directory) / "events.jsonl"),
                JsonlWriter(Path(directory) / "samples.jsonl"),
                calibration_data=calibration,
            )

            row = logger._make_event("click", 104, 208, source="test")

            self.assertEqual(row["region_x"], 4)
            self.assertEqual(row["region_y"], 8)
            self.assertEqual(row["video_x"], 4)
            self.assertEqual(row["video_y"], 8)
            self.assertFalse(row["calibration_applied"])
            self.assertEqual(row["calibration_method"], "raw_video_region")
            self.assertTrue(row["coordinate_check_completed"])

    def test_coordinate_check_mapping_is_diagnostic_only(self) -> None:
        with TemporaryDirectory() as directory:
            region = Region(screen_x=100, screen_y=200, width=400, height=300)
            calibration = build_calibration_result(
                region,
                self._scaled_calibration_clicks(region, scale_x=1.1, scale_y=0.9, shift_x=15, shift_y=-8),
                click_tolerance_px=80,
            )
            timing = TimingContext(session_id="test", logger_start_monotonic_ms=1000, video_zero_monotonic_ms=1000)
            logger = MouseActivityLogger(
                region,
                timing,
                AppConfig(),
                JsonlWriter(Path(directory) / "events.jsonl"),
                JsonlWriter(Path(directory) / "samples.jsonl"),
                calibration_data=calibration,
            )

            row = logger._make_event("click", 335, 327, source="test")

            self.assertTrue(calibration["completed"])
            self.assertEqual(calibration["mapping"]["method"], "visual_affine_least_squares")
            self.assertFalse(calibration["mapping"]["applied_to_recording_rows"])
            self.assertEqual(row["video_x"], 235)
            self.assertEqual(row["video_y"], 127)
            self.assertEqual(row["calibration_method"], "raw_video_region")

    def test_outside_filter_uses_raw_video_coordinates(self) -> None:
        with TemporaryDirectory() as directory:
            region = Region(screen_x=100, screen_y=100, width=100, height=100)
            calibration = build_calibration_result(region, self._calibration_clicks(region, offset_x=5, offset_y=0), 80)
            config = AppConfig(record_outside_region=False)
            timing = TimingContext(session_id="test", logger_start_monotonic_ms=1000, video_zero_monotonic_ms=1000)
            logger = MouseActivityLogger(
                region,
                timing,
                config,
                JsonlWriter(Path(directory) / "events.jsonl"),
                JsonlWriter(Path(directory) / "samples.jsonl"),
                calibration_data=calibration,
            )

            row = logger._make_event("click", 204, 150, source="test")
            logger._enqueue_event(row)

            self.assertFalse(row["inside_region"])
            self.assertFalse(row["inside_video_region"])
            self.assertEqual(logger.event_queue.qsize(), 0)

    def test_logger_ids_continue_across_segments(self) -> None:
        with TemporaryDirectory() as directory:
            region = Region(screen_x=0, screen_y=0, width=100, height=100)
            timing = TimingContext(session_id="test", logger_start_monotonic_ms=1000, video_zero_monotonic_ms=1000)
            first = MouseActivityLogger(
                region,
                timing,
                AppConfig(),
                JsonlWriter(Path(directory) / "events_1.jsonl"),
                JsonlWriter(Path(directory) / "samples_1.jsonl"),
            )
            first_event = first._make_event("click", 10, 10, source="test")
            first_sample = first._make_sample(10, 10, t_ms=1000)
            second = MouseActivityLogger(
                region,
                timing,
                AppConfig(),
                JsonlWriter(Path(directory) / "events_2.jsonl"),
                JsonlWriter(Path(directory) / "samples_2.jsonl"),
                event_counter_start=first.event_counter,
                sample_counter_start=first.sample_counter,
            )
            second_event = second._make_event("click", 20, 20, source="test")
            second_sample = second._make_sample(20, 20, t_ms=1100)

            self.assertEqual(first_event["event_id"], "evt_000001")
            self.assertEqual(first_sample["sample_id"], "smp_000001")
            self.assertEqual(second_event["event_id"], "evt_000002")
            self.assertEqual(second_sample["sample_id"], "smp_000002")

    def test_summary_and_xlsx(self) -> None:
        with TemporaryDirectory() as directory:
            storage = self._sample_storage(Path(directory))
            summary = generate_summary(storage)

            self.assertEqual(summary["clicks_total"], 1)
            self.assertEqual(summary["drag_count"], 1)
            self.assertTrue(storage.mouse_summary.exists())
            self.assertTrue(zipfile.is_zipfile(storage.mouse_summary_xlsx))
            self.assertTrue(zipfile.is_zipfile(storage.mouse_analysis_xlsx))
            with zipfile.ZipFile(storage.mouse_summary_xlsx) as archive:
                sheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
            self.assertIn("Mouse Events", sheet)
            self.assertIn("Mouse Samples", sheet)
            self.assertIn("video_x", sheet)
            with zipfile.ZipFile(storage.mouse_analysis_xlsx) as archive:
                analysis_sheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
            self.assertIn("录制概览", analysis_sheet)
            self.assertIn("操作明细", analysis_sheet)
            self.assertIn("视频X", analysis_sheet)

    def test_summary_prefers_video_region_flag(self) -> None:
        with TemporaryDirectory() as directory:
            storage = SessionStorage(Path(directory))
            storage.mouse_events.write_text(
                json.dumps(
                    {
                        "event_type": "click",
                        "t_video_ms": 10,
                        "inside_region": False,
                        "inside_video_region": True,
                    }
                )
                + "\n",
                encoding="utf-8",
            )
            storage.mouse_samples.write_text("", encoding="utf-8")

            summary = generate_summary(storage)

            self.assertEqual(summary["clicks_inside_region"], 1)
            self.assertEqual(summary["clicks_outside_region"], 0)

    def test_cli_postprocess(self) -> None:
        with TemporaryDirectory() as directory:
            storage = self._sample_storage(Path(directory))

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(ROOT / "run.py"),
                    "postprocess",
                    str(storage.session_dir),
                ],
                cwd=ROOT,
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn('"clicks_total": 1', result.stdout)
            self.assertTrue(storage.mouse_summary.exists())

    def test_behavior_report_from_session_folder(self) -> None:
        with TemporaryDirectory() as directory:
            storage = self._analysis_storage(Path(directory))

            result = generate_behavior_report(storage.session_dir)

            self.assertEqual(result.output_dir, (storage.session_dir / "analysis_output").resolve())
            self.assertEqual(result.metrics["clicks_total"], 2)
            self.assertTrue(zipfile.is_zipfile(result.outputs["report"]))
            self.assertTrue(result.outputs["heatmap_circle"].exists())
            self.assertFalse((result.output_dir / "click_heatmap_true_ratio.png").exists())
            self.assertFalse((result.output_dir / "click_heatmap_square_matrix.png").exists())
            workbook = load_workbook(result.outputs["report"], read_only=True)
            heatmap_sheet = workbook["点击热力图"]
            self.assertEqual(heatmap_sheet.max_row, 81)
            self.assertEqual(heatmap_sheet.max_column, 41)
            self.assertEqual(heatmap_sheet.cell(row=1, column=41).value, "列40")
            workbook.close()

    def test_behavior_report_from_summary_xlsx(self) -> None:
        with TemporaryDirectory() as directory:
            storage = self._analysis_storage(Path(directory))
            generate_summary(storage)

            result = generate_behavior_report(storage.mouse_summary_xlsx)

            self.assertEqual(result.output_dir, (storage.session_dir / "analysis_output").resolve())
            self.assertEqual(result.metrics["clicks_total"], 2)
            self.assertTrue(zipfile.is_zipfile(result.outputs["report"]))

    @staticmethod
    def _sample_storage(session_dir: Path) -> SessionStorage:
        storage = SessionStorage(session_dir)
        storage.mouse_events.write_text(
            "\n".join(
                [
                    json.dumps({"event_type": "left_down", "t_video_ms": 10, "inside_region": True}),
                    json.dumps({"event_type": "left_up", "t_video_ms": 80, "inside_region": True}),
                    json.dumps({"event_type": "click", "t_video_ms": 80, "inside_region": True}),
                    json.dumps({"event_type": "wheel", "t_video_ms": 120, "inside_region": False}),
                    json.dumps({"event_type": "drag_start", "t_video_ms": 200, "inside_region": True}),
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        storage.mouse_samples.write_text(
            "\n".join(
                [
                    json.dumps({"event_type": "move_sample", "t_video_ms": 0}),
                    json.dumps({"event_type": "move_sample", "t_video_ms": 1000}),
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        return storage

    @staticmethod
    def _analysis_storage(session_dir: Path) -> SessionStorage:
        storage = SessionStorage(session_dir)
        storage.session_meta.write_text(
            json.dumps({"recording_region": {"screen_x": 0, "screen_y": 0, "width": 100, "height": 200}}),
            encoding="utf-8",
        )
        storage.mouse_events.write_text(
            "\n".join(
                [
                    json.dumps(
                        {
                            "event_id": "evt_000001",
                            "event_type": "click",
                            "t_video_ms": 100,
                            "video_timecode": "00:00.100",
                            "video_x": 20,
                            "video_y": 40,
                            "inside_video_region": True,
                        }
                    ),
                    json.dumps(
                        {
                            "event_id": "evt_000002",
                            "event_type": "click",
                            "t_video_ms": 500,
                            "video_timecode": "00:00.500",
                            "video_x": 50,
                            "video_y": 160,
                            "inside_video_region": True,
                        }
                    ),
                    json.dumps(
                        {
                            "event_id": "evt_000003",
                            "event_type": "double_click_candidate",
                            "t_video_ms": 520,
                            "video_timecode": "00:00.520",
                            "video_x": 50,
                            "video_y": 160,
                            "inside_video_region": True,
                        }
                    ),
                    json.dumps(
                        {
                            "event_id": "evt_000004",
                            "event_type": "drag_start",
                            "t_video_ms": 1000,
                            "video_timecode": "00:01.000",
                            "video_x": 10,
                            "video_y": 20,
                            "inside_video_region": True,
                        }
                    ),
                    json.dumps(
                        {
                            "event_id": "evt_000005",
                            "event_type": "drag_end",
                            "t_video_ms": 1300,
                            "video_timecode": "00:01.300",
                            "video_x": 70,
                            "video_y": 80,
                            "inside_video_region": True,
                        }
                    ),
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        storage.mouse_samples.write_text(
            "\n".join(
                [
                    json.dumps({"sample_id": "smp_000001", "t_video_ms": 0, "video_x": 10, "video_y": 20}),
                    json.dumps({"sample_id": "smp_000002", "t_video_ms": 2000, "video_x": 20, "video_y": 30}),
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        return storage

    @staticmethod
    def _calibration_clicks(
        region: Region,
        offset_x: int,
        offset_y: int,
        event_offset_x: int = 0,
        event_offset_y: int = 0,
    ) -> list[dict[str, int | str]]:
        inset = 10
        targets = [
            ("top_left", region.screen_x + inset, region.screen_y + inset),
            ("top_right", region.screen_x + region.width - inset, region.screen_y + inset),
            ("bottom_left", region.screen_x + inset, region.screen_y + region.height - inset),
            ("bottom_right", region.screen_x + region.width - inset, region.screen_y + region.height - inset),
            ("center", region.screen_x + region.width // 2, region.screen_y + region.height // 2),
        ]
        return [
            {
                "target_id": target_id,
                "label": target_id,
                "expected_screen_x": x,
                "expected_screen_y": y,
                "actual_screen_x": x + offset_x,
                "actual_screen_y": y + offset_y,
                "tk_event_screen_x": x + event_offset_x,
                "tk_event_screen_y": y + event_offset_y,
            }
            for target_id, x, y in targets
        ]

    @staticmethod
    def _scaled_calibration_clicks(
        region: Region,
        scale_x: float,
        scale_y: float,
        shift_x: float,
        shift_y: float,
    ) -> list[dict[str, int | str]]:
        inset = 10
        targets = [
            ("top_left", inset, inset),
            ("top_right", region.width - inset, inset),
            ("bottom_left", inset, region.height - inset),
            ("bottom_right", region.width - inset, region.height - inset),
            ("center", region.width // 2, region.height // 2),
        ]
        rows = []
        for target_id, video_x, video_y in targets:
            expected_screen_x = region.screen_x + video_x
            expected_screen_y = region.screen_y + video_y
            actual_screen_x = int(round(region.screen_x + video_x * scale_x + shift_x))
            actual_screen_y = int(round(region.screen_y + video_y * scale_y + shift_y))
            rows.append(
                {
                    "target_id": target_id,
                    "label": target_id,
                    "expected_screen_x": expected_screen_x,
                    "expected_screen_y": expected_screen_y,
                    "actual_screen_x": actual_screen_x,
                    "actual_screen_y": actual_screen_y,
                    "tk_event_screen_x": expected_screen_x,
                    "tk_event_screen_y": expected_screen_y,
                }
            )
        return rows


if __name__ == "__main__":
    unittest.main()
